import csv
import io

from openpyxl import load_workbook

from app.extensions import db
from app.models import Employee
from tests.test_employees import _facility, create_employee


def test_view_page_is_read_only_and_downloads_one_employee(admin):
    employee = create_employee(admin, full_name="Viewable Person")
    page = admin.get(f"/employees/{employee.id}/view")
    assert page.status_code == 200
    body = page.get_data(as_text=True)
    assert "Viewable Person" in body
    assert "Save personal details" not in body
    assert "Download Excel" in body
    workbook = admin.get(f"/employees/export/excel?employee_id={employee.id}")
    assert workbook.status_code == 200
    assert "spreadsheetml" in workbook.content_type
    sheet = load_workbook(io.BytesIO(workbook.data)).active
    assert sheet["A1"].value == "Employee code"
    assert any(cell.value == "Viewable Person" for row in sheet.iter_rows() for cell in row)
    pdf = admin.get(f"/employees/export/pdf?employee_id={employee.id}")
    assert pdf.status_code == 200
    assert pdf.data.startswith(b"%PDF-")
    assert b"Viewable Person" in pdf.data


def test_register_download_follows_the_name_filter(admin):
    create_employee(admin, employee_code="EMP-3101", full_name="Kept Person")
    create_employee(admin, employee_code="EMP-3102", full_name="Other Person")
    workbook = admin.get("/employees/export/excel?q=Kept")
    sheet = load_workbook(io.BytesIO(workbook.data)).active
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "Kept Person" in values
    assert "Other Person" not in values
    listing = admin.get("/employees/?q=Kept")
    assert b"/view" in listing.data
    assert b"Bulk upload" in listing.data


def test_sample_csv_is_ignored_and_a_valid_row_is_added(admin):
    facility = _facility()
    template = admin.get("/employees/bulk/template")
    assert template.status_code == 200
    text = template.data.decode("utf-8-sig")
    assert "full_name" in text
    assert "Dr. Sample Name" in text
    rows = list(csv.reader(io.StringIO(text)))
    header = rows[0]
    sample = dict(zip(header, rows[1]))
    sample["full_name"] = "Bulk Person"
    sample["date_of_birth"] = "15/01/1990"
    sample["phone"] = "9867000444"
    sample["facility"] = facility.name
    payload = _csv([header, [sample[column] for column in header]])
    uploaded = admin.post("/employees/bulk", data={"file": (io.BytesIO(payload), "employees.csv")}, content_type="multipart/form-data")
    assert uploaded.status_code == 302
    stored = db.session.scalar(db.select(Employee).where(Employee.full_name == "Bulk Person"))
    assert stored is not None
    assert stored.phone == "9867000444"
    assert stored.employee_code.startswith("EMP-")
    sample_only = admin.post(
        "/employees/bulk",
        data={"file": (io.BytesIO(template.data), "employee-template.csv")},
        content_type="multipart/form-data",
    )
    assert sample_only.status_code == 400
    assert b"no employee rows" in sample_only.data


def test_invalid_bulk_row_adds_nobody(admin):
    facility = _facility()
    before = db.session.scalar(db.select(db.func.count()).select_from(Employee))
    header = ["full_name", "date_of_birth", "gender", "phone", "facility", "email", "status"]
    row = ["Broken Person", "1991-02-02", "female", "1234567890", facility.name, "", "active"]
    payload = _csv([header, row])
    response = admin.post("/employees/bulk", data={"file": (io.BytesIO(payload), "employees.csv")}, content_type="multipart/form-data")
    assert response.status_code == 400
    assert b"No employees were added" in response.data
    after = db.session.scalar(db.select(db.func.count()).select_from(Employee))
    assert after == before


def test_downloads_require_sign_in(client):
    assert client.get("/employees/export/excel").status_code == 302
    assert client.get("/employees/export/pdf").status_code == 302
    assert client.get("/employees/bulk").status_code == 302


def _csv(rows) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")
