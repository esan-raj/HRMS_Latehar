import csv
import io

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.extensions import db
from app.models import Department, Facility


def test_each_master_list_offers_bulk_download_and_upload(admin):
    for slug in ("departments", "designations", "facilities", "employment-types"):
        page = admin.get(f"/masters/{slug}")
        assert page.status_code == 200
        body = page.get_data(as_text=True)
        assert "Bulk upload" in body
        assert "Download Excel" in body
        assert "Download PDF" in body
        assert admin.get(f"/masters/{slug}/bulk").status_code == 200
        template = admin.get(f"/masters/{slug}/bulk/template")
        assert template.status_code == 200
        assert template.data.startswith(b"\xef\xbb\xbf")


def test_department_download_and_csv_upload(admin):
    created = admin.post("/masters/departments", data={"name": "Accounts", "version": "0"})
    assert created.status_code == 302
    workbook = admin.get("/masters/departments/export/excel")
    sheet = load_workbook(io.BytesIO(workbook.data)).active
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert "Accounts" in values
    pdf = admin.get("/masters/departments/export/pdf")
    assert pdf.data.startswith(b"%PDF-")
    assert b"Accounts" in pdf.data
    template = admin.get("/masters/departments/bulk/template").data.decode("utf-8-sig")
    header = next(csv.reader(io.StringIO(template)))
    payload = _csv([header, ["Ward office", "active"]])
    uploaded = admin.post(
        "/masters/departments/bulk",
        data={"file": (io.BytesIO(payload), "departments.csv")},
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 302
    stored = db.session.scalar(select(Department).where(Department.name == "Ward office"))
    assert stored is not None
    assert stored.is_active is True
    sample = admin.post(
        "/masters/departments/bulk",
        data={"file": (io.BytesIO(admin.get("/masters/departments/bulk/template").data), "departments.csv")},
        content_type="multipart/form-data",
    )
    assert sample.status_code == 400
    assert b"no rows to add" in sample.data


def test_duplicate_master_row_adds_nothing(admin):
    payload = _csv([["name", "status"], ["Accounts", "active"], ["Accounts", "inactive"]])
    response = admin.post(
        "/masters/departments/bulk",
        data={"file": (io.BytesIO(payload), "departments.csv")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"Nothing was added" in response.data
    count = db.session.scalar(select(func.count()).select_from(Department))
    assert count == 0


def test_facility_code_is_assigned_and_cannot_be_changed(admin):
    page = admin.get("/masters/facilities")
    assert b'name="code"' not in page.data
    assert b"FAC-0001" in page.data
    first = admin.post(
        "/masters/facilities",
        data={"name": "Community Health Centre", "facility_type": "CHC", "district": "Latehar", "version": "0"},
    )
    assert first.status_code == 302, first.get_data(as_text=True)
    second = admin.post(
        "/masters/facilities",
        data={"name": "Referral Hospital", "facility_type": "Hospital", "district": "Latehar", "block": "Garu", "version": "0"},
    )
    assert second.status_code == 302
    stored = db.session.scalars(select(Facility).order_by(Facility.code)).all()
    assert [row.code for row in stored] == ["FAC-0001", "FAC-0002"]
    edited = admin.post(
        f"/masters/facilities/{stored[0].id}/edit",
        data={
            "name": "Community Health Centre",
            "code": "FAC-9999",
            "facility_type": "CHC",
            "district": "Latehar",
            "version": "1",
        },
    )
    assert edited.status_code == 302
    assert db.session.get(Facility, stored[0].id).code == "FAC-0001"


def test_facility_csv_receives_the_next_code(admin):
    payload = _csv(
        [
            ["name", "facility_type", "district", "block", "status"],
            ["Community Health Centre", "CHC", "Latehar", "", "active"],
        ]
    )
    uploaded = admin.post(
        "/masters/facilities/bulk",
        data={"file": (io.BytesIO(payload), "facilities.csv")},
        content_type="multipart/form-data",
    )
    assert uploaded.status_code == 302, uploaded.get_data(as_text=True)
    stored = db.session.scalar(select(Facility).where(Facility.name == "Community Health Centre"))
    assert stored is not None
    assert stored.code == "FAC-0001"
    assert stored.district == "Latehar"
    assert stored.block is None


def test_master_downloads_require_sign_in(client):
    assert client.get("/masters/designations/export/excel").status_code == 302
    assert client.get("/masters/facilities/bulk").status_code == 302


def _csv(rows) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")
